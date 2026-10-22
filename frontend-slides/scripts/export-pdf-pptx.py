"""HTML (frontend-slides) -> PDF + PPTX, sem Chromium/LibreOffice.
Slides com data-steps=k (cliques) viram k+1 paginas no PDF/PPTX.
Uso: python3 export-pdf-pptx.py <index.html> <pasta_temporaria_para_PNGs>
Gera <index>.pdf e <index>.pptx ao lado do HTML (um PNG 1920x1080 por slide;
o texto de cada slide vai para as notas do PPTX). Navega por #N e ?static.
Navegador: variável FIREFOX, ou 'firefox' no PATH, ou o do cluster (gridunesp).
"""
import subprocess, sys, re, os, shutil
from html.parser import HTMLParser
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
SRC=os.path.abspath(sys.argv[1]); OUT=os.path.splitext(SRC)[0]; SHOTS=sys.argv[2]
FF=os.environ.get('FIREFOX') or shutil.which('firefox') or '/opt/gridunesp/dist/firefox/74.0/firefox/firefox'
os.makedirs(SHOTS,exist_ok=True)
html=open(SRC,encoding='utf-8').read()
n=html.count('<section class="slide')
class T(HTMLParser):
    def __init__(s): super().__init__(); s.sl=[]; s.st=[]; s.cur=None; s.skip=0
    def handle_starttag(s,t,a):
        if t=='section': s.cur=[]; s.sl.append(s.cur); s.st.append(int(dict(a).get('data-steps',0)))
        if t in('script','style'): s.skip+=1
    def handle_endtag(s,t):
        if t in('script','style'): s.skip-=1
        if t in('h1','h2','li','p','tr','div') and s.cur is not None: s.cur.append('\n')
    def handle_data(s,d):
        if s.cur is not None and not s.skip: s.cur.append(d)
p=T(); p.feed(html)
notes=[re.sub(r'\n\s*\n+','\n',''.join(x)).strip() for x in p.sl]
imgs=[]; pages=[]  # slide com data-steps=k gera k+1 paginas (estado inicial + um por clique)
for i in range(1,n+1):
    for k in range(p.st[i-1]+1):
        f=f'{SHOTS}/s{i:02d}'+(f'_{k}' if k else '')+'.png'
        subprocess.run([FF,'--headless','--screenshot',f,'--window-size=1920,1080',f'file://{SRC}?static'+(f'&step={k}' if k else '')+f'#{i}'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=120,check=True)
        im=Image.open(f); assert im.size==(1920,1080),im.size; imgs.append(f); pages.append(notes[i-1])
notes=pages
# PDF
pil=[Image.open(f).convert('RGB') for f in imgs]
pil[0].save(OUT+'.pdf',save_all=True,append_images=pil[1:],resolution=144)
# PPTX 16:9 13.333x7.5in
pr=Presentation(); pr.slide_width=Emu(12192000); pr.slide_height=Emu(6858000)
for f,nt in zip(imgs,notes):
    s=pr.slides.add_slide(pr.slide_layouts[6])
    s.shapes.add_picture(f,0,0,pr.slide_width,pr.slide_height)
    s.notes_slide.notes_text_frame.text=nt
pr.save(OUT+'.pptx')
print(n,'slides')
